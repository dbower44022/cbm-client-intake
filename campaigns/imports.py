"""The imports (build plan § 2, reworked for the lead model § 7a): parse →
normalise → already-known check → plan → apply. Three Lead Sources share every
step: the state filing (§ 2.1), the purchased list and the partner file (§ 2.3).

**Every row becomes a lead, and nothing else is created** (rulings 41, 43): the
Company and the Contact are made at first contact, never at import.

**The status a lead enters at (ruling 53).** A state-filing row that carries a
website or an email enters **Qualified** on the outside filter's verdict
(ruling 11), with the filter recorded in ``qualifiedBy`` / ``qualifiedAt``.
Every other row enters **Unqualified** (rulings 18, 42).

**Plan, then apply.** The upload shows what would be created, matched and
skipped, with the counts, and nothing is written until the reviewer applies that
exact plan (the Operations-job pattern). The plan is stored in the ledger with
the normalised rows, so the apply writes what was reviewed.

**Who acts (ruling 33).** The already-known reads and the link to a matched
Company run under the intake API client; the leads are created as the
signed-in user. **A re-upload of the same file is a no-op**: the ledger knows
every row an applied batch created.

**The already-known check (ruling 12)** is conservative, close to exact: a
Company by name AND postal code, a Company by website key, a Contact by email. A
match is a **Duplicate** (ruling 50) — a converted lead's Company or Contact
included, since the import no longer reuses one at a later conversion — the lead is still created (it is the
record that the row arrived and was recognised), linked to the Company it
matched, or naming the Contact it matched in its reason note, and no touch is
ever planned for it.
"""

from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import logging
import re
from datetime import datetime, timedelta, timezone
from typing import Any, Awaitable, Callable, Optional

from core.config import Settings, get_settings
from core.crm_upsert import website_key
from core.espo import EspoApi, EspoError, is_not_found

from . import config as cfg
from . import import_templates as tpl
from .service import CampaignError
from .store import PLAN_FAILED, PLAN_QUEUED, PLAN_READY, PLAN_RUNNING

log = logging.getLogger("cbm_intake.campaigns.imports")

REQUIRED = ("business_name",)
ADDRESS_SET = ("mailing_street", "mailing_city", "mailing_state", "mailing_zip")
KNOWN_COLUMNS = REQUIRED + ADDRESS_SET + (
    "external_reference", "website", "email", "phone", "contact_first_name",
    "contact_last_name", "qualified_by", "list_source",
)
MAX_ROWS = 5000
MAX_TEXT = 5 * 1024 * 1024

OUTCOME_CREATE = "create"
OUTCOME_ALREADY_KNOWN = "already_known"
OUTCOME_DUPLICATE = "duplicate"          # the same row twice in one file
OUTCOME_KNOWN_ROW = "known_row"          # an applied batch already created it
OUTCOME_INVALID = "invalid"
OUTCOME_FILLS = "fills"                  # ruling 84: a repeat row fills a lead's empty fields (M4)

#: The row fields a repeat row may fill, and the lead attribute each one is
#: (ruling 84). The business name and the status are never filled: the name is
#: the row's identity, the status is the ladder's.
FILL_FIELDS: dict[str, str] = {
    "contact_first_name": "firstName", "contact_last_name": "lastName",
    "mailing_street": "mailingAddressStreet", "mailing_city": "mailingAddressCity",
    "mailing_state": "mailingAddressState", "mailing_zip": "mailingAddressPostalCode",
    "website": "website", "email": "email", "phone": "phone",
    "external_reference": "externalReference",
}
FILL_SELECT = "id,name,businessName," + ",".join(FILL_FIELDS.values())
#: The attribute names as the preview's note says them.
FILL_WORDS: dict[str, str] = {
    "firstName": "first name", "lastName": "last name", "mailingAddressStreet": "street",
    "mailingAddressCity": "city", "mailingAddressState": "state", "mailingAddressPostalCode": "ZIP",
    "website": "website", "email": "email", "phone": "phone", "externalReference": "reference",
}
_ID_CHUNK = 100

#: Undo (ruling 87): a change inside this long after the row was written is
#: the import's own (the create, the Duplicate link, placement), not a touch.
UNDO_WINDOW = timedelta(minutes=5)
UNDO_SELECT = ("id,name,businessName,leadStatus,leadStatusChangedAt,sideExit,sideExitAt,"
               "qualificationLevel,levelSetAt,modifiedAt,createdAt")
CONVERSATIONS_LINK = "conversations"     # CProspect.conversations (marketing-redesign-crm-handoff.md § 1)

_ZIP = re.compile(r"^\d{5}(-\d{4})?$")


def _now_crm() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")


def _norm_header(h: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", (h or "").strip().lower()).strip("_")


def parse_csv(text: str) -> tuple[list[str], list[dict[str, str]]]:
    """Columns (normalised) and rows. A missing required column fails the plan
    with the column named; unknown columns ride along into ``importRow``."""
    if len(text) > MAX_TEXT:
        raise CampaignError("The file is too large to import in one go (5 MB limit).")
    reader = csv.reader(io.StringIO(text.lstrip("﻿")))
    try:
        header = next(reader)
    except StopIteration:
        raise CampaignError("The file is empty.")
    columns = [_norm_header(h) for h in header]
    missing = [c for c in REQUIRED if c not in columns]
    if missing:
        raise CampaignError(f"The file has no {', '.join(missing)} column.")
    if not any(c in columns for c in ADDRESS_SET):
        raise CampaignError("The file has no mailing address columns "
                            "(mailing_street, mailing_city, mailing_state, mailing_zip).")
    rows: list[dict[str, str]] = []
    for raw in reader:
        if not any((v or "").strip() for v in raw):
            continue
        row = {columns[i]: (raw[i].strip() if i < len(raw) else "") for i in range(len(columns)) if columns[i]}
        rows.append(row)
        if len(rows) > MAX_ROWS:
            raise CampaignError(f"The file has more than {MAX_ROWS} rows; split it.")
    return columns, rows


def normalise(row: dict[str, str]) -> dict[str, Any]:
    out = dict(row)
    out["business_name"] = re.sub(r"\s+", " ", row.get("business_name", "")).strip()
    out["mailing_state"] = (row.get("mailing_state") or "").strip().upper()[:2]
    zip_ = (row.get("mailing_zip") or "").strip()
    if len(zip_) == 9 and zip_.isdigit():
        zip_ = f"{zip_[:5]}-{zip_[5:]}"
    out["mailing_zip"] = zip_
    out["email"] = (row.get("email") or "").strip().lower()
    out["website"] = (row.get("website") or "").strip()
    out["external_reference"] = (row.get("external_reference") or "").strip()
    return out


def row_hash(door: str, row: dict[str, Any]) -> str:
    """The across-batch key (§ 2.1): the external reference when the row has
    one, else business name plus postal code."""
    ref = row.get("external_reference") or ""
    if ref:
        key = f"{door}|ref|{ref.lower()}"
    else:
        key = f"{door}|name|{(row.get('business_name') or '').lower()}|{(row.get('mailing_zip') or '')[:5]}"
    return hashlib.sha256(key.encode()).hexdigest()


def file_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8", "replace")).hexdigest()


def bytes_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


# --- reading a file: CSV or Excel (ruling 80; M4) -----------------------------------------------

class SheetChoice(CampaignError):
    """A workbook with several sheets: the page asks which, then posts again
    with ``sheet``."""

    def __init__(self, sheets: list[str]):
        super().__init__("Which sheet? The workbook has several: " + ", ".join(sheets) + ".")
        self.sheets = sheets


def decode_file(file_base64: str) -> bytes:
    """The file as posted — base64 inside the JSON, the way the CSV text has
    always travelled; never multipart."""
    try:
        data = base64.b64decode(file_base64 or "", validate=False)
    except Exception:  # noqa: BLE001 — a bad body is the person's to fix
        raise CampaignError("The file could not be read from the upload.")
    if len(data) > MAX_TEXT:
        raise CampaignError("The file is too large to import in one go (5 MB limit).")
    return data


def _table_rows(headers_raw: list[Any], body: list[list[Any]]) -> tuple[list[str], list[dict[str, str]]]:
    columns = [tpl.norm_header(h) for h in headers_raw]
    if not any(columns):
        raise CampaignError("The file's first row has no column headings.")
    rows: list[dict[str, str]] = []
    for raw in body:
        cells = ["" if v is None else str(v).strip() for v in raw]
        if not any(cells):
            continue
        rows.append({columns[i]: (cells[i] if i < len(cells) else "") for i in range(len(columns)) if columns[i]})
        if len(rows) > MAX_ROWS:
            raise CampaignError(f"The file has more than {MAX_ROWS} rows; split it.")
    return columns, rows


def _read_csv_table(text: str) -> tuple[list[str], list[dict[str, str]]]:
    if len(text) > MAX_TEXT:
        raise CampaignError("The file is too large to import in one go (5 MB limit).")
    reader = csv.reader(io.StringIO(text.lstrip("﻿")))
    try:
        header = next(reader)
    except StopIteration:
        raise CampaignError("The file is empty.")
    return _table_rows(header, list(reader))


def _cell_text(value: Any) -> str:
    """A spreadsheet cell as the text a person sees: whole numbers without a
    trailing ``.0`` (a ZIP, a phone, a reference), dates as ISO."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat() if (value.hour, value.minute, value.second) == (0, 0, 0) else value.isoformat(sep=" ")
    return str(value).strip()


def _read_xlsx(data: bytes, sheet: Optional[str]) -> tuple[list[str], list[dict[str, str]], str]:
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover — pinned in pyproject.toml
        raise CampaignError("This deployment cannot read Excel files (openpyxl is missing).")
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — not a workbook
        raise CampaignError("The file is not an Excel workbook (.xlsx) that can be opened.")
    try:
        names = [ws.title for ws in book.worksheets]
        if not names:
            raise CampaignError("The workbook has no sheets.")
        if sheet:
            if sheet not in names:
                raise CampaignError(f"The workbook has no sheet named {sheet!r}: " + ", ".join(names) + ".")
            chosen = sheet
        elif len(names) == 1:
            chosen = names[0]
        else:
            raise SheetChoice(names)
        ws = book[chosen]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = [_cell_text(v) for v in next(rows_iter)]
        except StopIteration:
            raise CampaignError(f"The sheet {chosen!r} is empty.")
        body = [[_cell_text(v) for v in r] for r in rows_iter]
    finally:
        book.close()
    columns, rows = _table_rows(header, body)
    return columns, rows, chosen


def read_file(name: str, data: bytes | str, sheet: Optional[str] = None) -> dict[str, Any]:
    """A CSV or ``.xlsx`` file as ``{"columns": [...normalised headers],
    "rows": [...], "sheet": name or None}``. ``.xls`` is refused with "save it
    as .xlsx"; a workbook with several sheets raises :class:`SheetChoice`."""
    lower = (name or "").lower()
    if lower.endswith(".xls"):
        raise CampaignError("Old Excel files (.xls) cannot be read — open it in Excel and save it as .xlsx.")
    if lower.endswith((".xlsx", ".xlsm")):
        if isinstance(data, str):
            data = data.encode("utf-8", "replace")
        columns, rows, chosen = _read_xlsx(data, sheet)
        return {"columns": columns, "rows": rows, "sheet": chosen}
    if isinstance(data, bytes):
        if data[:4] == b"PK\x03\x04":
            columns, rows, chosen = _read_xlsx(data, sheet)
            return {"columns": columns, "rows": rows, "sheet": chosen}
        text = data.decode("utf-8-sig", "replace")
    else:
        text = data
    columns, rows = _read_csv_table(text)
    return {"columns": columns, "rows": rows, "sheet": None}


def _invalid_reason(row: dict[str, Any]) -> Optional[str]:
    if not row.get("business_name"):
        return "no business name"
    if not all(row.get(c) for c in ADDRESS_SET):
        return "incomplete mailing address"
    if not _ZIP.match(row.get("mailing_zip") or ""):
        return "postal code is not a US ZIP"
    return None


# --- the already-known check (ruling 12) -----------------------------------------------

async def already_known(api_client: EspoApi, row: dict[str, Any]) -> Optional[dict[str, Any]]:
    """The Company (or Contact) this row already is, or None. Conservative and
    close to exact; each step is best-effort so one unreadable list never blocks
    an import (a miss errs toward mailing a card — ruling 12)."""
    name = row.get("business_name") or ""
    zip5 = (row.get("mailing_zip") or "")[:5]
    if name and zip5:
        try:
            page = await api_client.list(
                cfg.ACCOUNT, select="id,name,billingAddressPostalCode,website,cCompanyType", max_size=20,
                where=[{"type": "equals", "attribute": "name", "value": name}],
            )
            for c in page.get("list", []):
                if ((c.get("billingAddressPostalCode") or "")[:5]) == zip5:
                    return {"kind": "company", "id": c["id"], "name": c.get("name"), "by": "name and postal code"}
        except EspoError as exc:
            log.warning("already-known name check failed for %r: %s", name, exc)
    if row.get("website"):
        try:
            c = await company_by_website(api_client, row["website"])
            if c:
                return {"kind": "company", "id": c["id"], "name": c.get("name"), "by": "website"}
        except EspoError as exc:
            log.warning("already-known website check failed for %r: %s", row["website"], exc)
    if row.get("email"):
        try:
            found = await api_client.find_one(cfg.CONTACT, "emailAddress", row["email"], select="id,name,cContactType")
            if found:
                return {"kind": "contact", "id": found["id"], "name": found.get("name"), "by": "email"}
        except EspoError as exc:
            log.warning("already-known email check failed for %r: %s", row["email"], exc)
    return None


async def company_by_website(api_client: EspoApi, website: str) -> Optional[dict[str, Any]]:
    """The Company at this web address, by website key (scheme, www., case and a
    trailing slash ignored; the path counts), or None."""
    key = website_key(website)
    if not key:
        return None
    host = key.split("/")[0]
    page = await api_client.list(
        cfg.ACCOUNT, select="id,name,website,cCompanyType", max_size=50,
        where=[{"type": "contains", "attribute": "website", "value": host}],
    )
    for c in page.get("list", []):
        if website_key(c.get("website")) == key:
            return c
    return None


_CHUNK = 40


class KnownIndex:
    """The already-known check (ruling 12) for a whole file in a few CRM
    reads instead of one to three per row — a month's filing is about 800
    rows (build plan § 7b), and row by row the plan took 61 seconds and timed
    the browser out (10-20-26). Same three matches, same order, same
    best-effort: a failed read is logged and that check simply finds nothing."""

    def __init__(self) -> None:
        self.by_name: dict[str, list[dict[str, Any]]] = {}
        self.by_website: dict[str, dict[str, Any]] = {}
        self.by_email: dict[str, dict[str, Any]] = {}

    @classmethod
    async def build(cls, api_client: EspoApi, rows: list[dict[str, Any]]) -> "KnownIndex":
        idx = cls()
        names = sorted({r["business_name"] for r in rows if r.get("business_name")})
        for i in range(0, len(names), _CHUNK):
            chunk = names[i:i + _CHUNK]
            try:
                page = await api_client.list(
                    cfg.ACCOUNT, select="id,name,billingAddressPostalCode,website,cCompanyType", max_size=200,
                    where=[{"type": "in", "attribute": "name", "value": chunk}])
                for c in page.get("list", []):
                    idx.by_name.setdefault((c.get("name") or "").lower(), []).append(c)
            except EspoError as exc:
                log.warning("already-known name check failed for %s names: %s", len(chunk), exc)
        keys = sorted({k for k in (website_key(r.get("website")) for r in rows) if k})
        for i in range(0, len(keys), 20):
            chunk = keys[i:i + 20]
            try:
                page = await api_client.list(
                    cfg.ACCOUNT, select="id,name,website,cCompanyType", max_size=200,
                    where=[{"type": "or", "value": [{"type": "contains", "attribute": "website",
                                                     "value": k.split("/")[0]} for k in chunk]}])
                for c in page.get("list", []):
                    k = website_key(c.get("website"))
                    if k and k not in idx.by_website:
                        idx.by_website[k] = c
            except EspoError as exc:
                log.warning("already-known website check failed for %s addresses: %s", len(chunk), exc)
        emails = sorted({r["email"] for r in rows if r.get("email")})
        for i in range(0, len(emails), 20):
            chunk = emails[i:i + 20]
            try:
                page = await api_client.list(
                    cfg.CONTACT, select="id,name,emailAddress,cContactType", max_size=200,
                    where=[{"type": "or", "value": [{"type": "equals", "attribute": "emailAddress",
                                                     "value": e} for e in chunk]}])
                for c in page.get("list", []):
                    e = (c.get("emailAddress") or "").lower()
                    if e and e not in idx.by_email:
                        idx.by_email[e] = c
            except EspoError as exc:
                log.warning("already-known email check failed for %s addresses: %s", len(chunk), exc)
        return idx

    def match(self, row: dict[str, Any]) -> Optional[dict[str, Any]]:
        name = (row.get("business_name") or "").lower()
        zip5 = (row.get("mailing_zip") or "")[:5]
        if name and zip5:
            for c in self.by_name.get(name, []):
                if ((c.get("billingAddressPostalCode") or "")[:5]) == zip5:
                    return {"kind": "company", "id": c["id"], "name": c.get("name"), "by": "name and postal code"}
        k = website_key(row.get("website"))
        if k and k in self.by_website:
            c = self.by_website[k]
            return {"kind": "company", "id": c["id"], "name": c.get("name"), "by": "website"}
        e = (row.get("email") or "").lower()
        if e and e in self.by_email:
            c = self.by_email[e]
            return {"kind": "contact", "id": c["id"], "name": c.get("name"), "by": "email"}
        return None


# --- plan ----------------------------------------------------------------------------------

Progress = Callable[[int, int], Awaitable[None]]


def _empty_counts(n: int) -> dict[str, int]:
    return {"rows": n, "create": 0, "alreadyKnown": 0, "duplicate": 0, "knownRow": 0, "invalid": 0,
            "fills": 0, "qualified": 0, "unqualified": 0,
            "identified": 0, "contactable": 0, "referred": 0}


async def _fill_candidates(api_client: EspoApi, store, rows: list[dict[str, Any]], hashes: list[str],
                           known: set[str]) -> dict[str, dict[str, Any]]:
    """Ruling 84: for every known row, the lead it created and the fields the
    row would fill on it — only the lead's EMPTY ones, read in id batches of
    100 (never a per-row call). ``{hash: {"leadId", "leadName", "fill"}}``."""
    wanted = {h: r for r, h in zip(rows, hashes) if h in known and not _invalid_reason(r)}
    if not wanted:
        return {}
    lead_by_hash = await store.leads_for_hashes(list(wanted))
    ids = sorted({i for i in lead_by_hash.values() if i})
    leads: dict[str, dict[str, Any]] = {}
    for i in range(0, len(ids), _ID_CHUNK):
        chunk = ids[i:i + _ID_CHUNK]
        try:
            page = await api_client.list(cfg.PROSPECT, select=FILL_SELECT, max_size=len(chunk),
                                         where=[{"type": "in", "attribute": "id", "value": chunk}])
            for lead in page.get("list", []):
                leads[lead["id"]] = lead
        except EspoError as exc:
            log.warning("import fills: %s leads could not be read: %s", len(chunk), exc)
    out: dict[str, dict[str, Any]] = {}
    for h, row in wanted.items():
        lead_id = lead_by_hash.get(h) or ""
        lead = leads.get(lead_id)
        if not lead:
            # The ledger names a lead the CRM no longer has (deleted by hand,
            # or by an Undo whose fills rows outlived it — round two of the M4
            # pass, 10-25-26): the row is new again, not a repeat.
            out[h] = {"leadId": None, "leadName": "", "fill": {}, "gone": True}
            continue
        fill = {attr: row[field] for field, attr in FILL_FIELDS.items()
                if row.get(field) and not str(lead.get(attr) or "").strip()}
        out[h] = {"leadId": lead["id"], "leadName": lead.get("businessName") or lead.get("name") or lead["id"],
                  "fill": fill}
    return out


def entry_level(row: dict[str, Any], settings: dict[str, Any]) -> str:
    """The level a Qualified row lands at (ruling 111): the ladder's own rule,
    with the template's partner as the referrer fact."""
    from .ladder import facts, placed
    probe = {"firstName": row.get("contact_first_name"), "lastName": row.get("contact_last_name"),
             "phone": row.get("phone"), "email": row.get("email"),
             "referredByPartnerId": settings.get("referrerPartnerId") or ""}
    return placed(facts(probe))


async def _plan_entries(api_client: EspoApi, store, rows: list[dict[str, Any]], settings: dict[str, Any],
                        *, templated: bool, ladder: bool, progress: Optional[Progress] = None
                        ) -> tuple[list[dict[str, Any]], dict[str, int]]:
    """Every row's outcome. ``templated`` is the M4 path (fills, Enters at, the
    level); without it a known row is skipped exactly as before."""
    door = settings["door"]
    hashes = [row_hash(door, r) for r in rows]
    known = await store.known_row_hashes(list(set(hashes)))
    fills = await _fill_candidates(api_client, store, rows, hashes, known) if templated else {}
    seen: set[str] = set()
    planned: list[dict[str, Any]] = []
    index = await KnownIndex.build(api_client, [r for r, h in zip(rows, hashes)
                                                if h not in known and not _invalid_reason(r)])
    counts = _empty_counts(len(rows))
    for n, (row, h) in enumerate(zip(rows, hashes), 1):
        entry: dict[str, Any] = {"row": row, "hash": h, "name": row.get("business_name") or "",
                                 "postalCode": row.get("mailing_zip") or ""}
        reason = _invalid_reason(row)
        if reason:
            entry.update(outcome=OUTCOME_INVALID, note=reason); counts["invalid"] += 1
        elif h in seen:
            entry.update(outcome=OUTCOME_DUPLICATE, note="the same row appears earlier in this file")
            counts["duplicate"] += 1
        elif h in known and not (fills.get(h) or {}).get("gone"):
            cand = fills.get(h)
            if cand and cand["fill"]:
                entry.update(outcome=OUTCOME_FILLS, leadId=cand["leadId"], fill=cand["fill"],
                             note=f"fills {', '.join(FILL_WORDS.get(k, k) for k in sorted(cand['fill']))} on {cand['leadName']}")
                counts["fills"] += 1
            else:
                entry.update(outcome=OUTCOME_KNOWN_ROW,
                             note=("an earlier import already created this row, and it has nothing to fill"
                                   if templated else "an earlier import already created this row"))
                if cand:
                    entry["leadId"] = cand["leadId"]
                counts["knownRow"] += 1
        else:
            match = index.match(row)
            if match:
                entry.update(outcome=OUTCOME_ALREADY_KNOWN, match=match,
                             note=f"matches {match['name'] or match['id']} by {match['by']}")
                counts["alreadyKnown"] += 1
            else:
                entry["outcome"] = OUTCOME_CREATE
                entry["status"] = entry_status_for(settings, row)
                if (fills.get(h) or {}).get("gone"):
                    entry["stale"] = True       # the ledger named a lead the CRM no longer has
                counts["create"] += 1
                counts["qualified" if entry["status"] == cfg.LEAD_QUALIFIED else "unqualified"] += 1
                if entry["status"] == cfg.LEAD_QUALIFIED and ladder:
                    entry["level"] = entry_level(row, settings)
                    counts[entry["level"].lower()] += 1
                elif templated and settings.get("entersAt") == cfg.LEAD_QUALIFIED:
                    entry["note"] = "enters Unqualified: no website or email"
        seen.add(h)
        planned.append(entry)
        if progress and n % 100 == 0:
            await progress(n, len(rows))
    return planned, counts


def _plan_answer(batch: dict[str, Any], plan: dict[str, Any]) -> dict[str, Any]:
    counts = plan.get("counts") or _empty_counts(0)
    rows = plan.get("rows") or []
    return {"batchId": batch["id"], "name": batch["name"], "counts": counts,
            "nothingToDo": counts.get("create", 0) + counts.get("alreadyKnown", 0) + counts.get("fills", 0) == 0,
            "priorBatch": plan.get("priorBatch"), "ladder": bool(plan.get("ladder")),
            "excludedRows": list(batch.get("excluded_rows") or []),
            "rows": [{k: v for k, v in e.items() if k != "row"} | {
                "website": e["row"].get("website") or "", "email": e["row"].get("email") or "",
                "row": e["row"]}
                     for e in rows[:500]]}


async def plan_import(api_client: EspoApi, store, *, text: str, name: str, door: str,
                      qualified_by: str, list_source: str = "") -> dict[str, Any]:
    """What an import would do (the CSV path, as it has always worked: the
    file's own column names, the plan computed inline, a known row skipped).
    Nothing is written; the plan is stored in the ledger as a planned batch and
    the apply names it."""
    if door not in cfg.IMPORT_SOURCES:
        raise CampaignError("Unknown Lead Source for an import.")
    sha = file_hash(text)
    prior = await store.find_batch_by_file(sha)
    columns, raw_rows = parse_csv(text)
    rows = [normalise(r) for r in raw_rows]
    settings = {"door": door, "qualifiedBy": qualified_by, "listSource": list_source}
    planned, counts = await _plan_entries(api_client, store, rows, settings, templated=False, ladder=False)
    plan = {
        "door": door, "qualifiedBy": qualified_by, "listSource": list_source,
        "columns": columns, "counts": counts, "rows": planned,
        "priorBatch": {"id": prior["id"], "appliedAt": prior["applied_at"].isoformat()} if prior else None,
    }
    batch = await store.create_batch(name=name, door=door, file_sha256=sha, row_count=len(rows),
                                     plan=plan, created_by="")
    answer = _plan_answer(batch, plan)
    for r in answer["rows"]:
        r.pop("row", None)
    return answer


# --- the templated import: queue, then the worker plans (rulings 83, 123; M4) -------------------

def mapped_rows(table: dict[str, Any], mapping: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """The file's rows through the template's mapping, normalised. A required
    field the mapping leaves out stops here with the field named (ruling 81)."""
    missing = tpl.missing_required(mapping)
    if missing:
        raise CampaignError("Nothing in the file is mapped onto " + ", ".join(missing) + ". Map those columns first.")
    return [normalise(tpl.apply_mapping(mapping, raw)) for raw in table["rows"]]


async def queue_plan(store, *, name: str, settings: dict[str, Any], table: dict[str, Any],
                     file_sha256: str, file_name: str, created_by: str) -> dict[str, Any]:
    """Create the batch at ``queued`` holding the mapped rows; the worker plans
    it (ruling 123) and the page polls. Nothing is read from the CRM here."""
    if settings["door"] not in cfg.IMPORT_SOURCES:
        raise CampaignError("Unknown Lead Source for an import.")
    rows = mapped_rows(table, settings.get("mapping") or [])
    if not rows:
        raise CampaignError("The file has no rows under its headings.")
    plan = {**{k: v for k, v in settings.items() if k != "mapping"}, "mapping": settings.get("mapping") or [],
            "columns": table["columns"], "sheet": table.get("sheet"), "input": rows}
    batch = await store.create_batch(name=name, door=settings["door"], file_sha256=file_sha256, row_count=len(rows),
                                     plan=plan, created_by=created_by, plan_status=PLAN_QUEUED,
                                     template_id=settings.get("templateId"), file_name=file_name)
    if settings.get("templateId"):
        try:
            await store.set_template_last_batch(settings["templateId"], batch["id"])
        except Exception as exc:  # noqa: BLE001 — bookkeeping only
            log.warning("template %s last batch not recorded: %s", settings["templateId"], exc)
    return batch


async def run_plan(api_client: EspoApi, store, batch: dict[str, Any], *, ladder: bool,
                   progress: Optional[Progress] = None) -> dict[str, Any]:
    """Plan a queued batch: the fourth outcome, the level per row (when the
    ladder is on and the CRM has it), the across-batch and in-file checks.
    Returns the finished plan; the caller stores it."""
    plan = dict(batch.get("plan") or {})
    rows = plan.pop("input", None) or []
    settings = {k: plan.get(k) for k in ("door", "qualifiedBy", "listSource", "entersAt",
                                        "referrerPartnerId", "referrerPartnerName", "templateId", "templateName")}
    settings["mapping"] = plan.get("mapping") or []
    prior = await store.find_batch_by_file(batch["file_sha256"])
    planned, counts = await _plan_entries(api_client, store, rows, settings, templated=True, ladder=ladder,
                                          progress=progress)
    plan.update(counts=counts, rows=planned, ladder=ladder,
                priorBatch=({"id": prior["id"], "appliedAt": prior["applied_at"].isoformat()}
                            if prior and prior["id"] != batch["id"] else None))
    return plan


async def plan_batch_now(api_client: EspoApi, store, batch_id: str, *, ladder: bool) -> dict[str, Any]:
    """Plan one queued batch to completion and store the result — the worker's
    unit of work, and the dry-run (dev) inline path."""
    batch = await store.get_batch(batch_id)
    if not batch:
        raise CampaignError("That import was not found.")

    async def progress(done: int, total: int) -> None:
        await store.set_plan_progress(batch_id, {"checked": done, "of": total,
                                                 "at": datetime.now(timezone.utc).isoformat()})

    try:
        await store.set_plan_progress(batch_id, {"checked": 0, "of": batch.get("row_count") or 0,
                                                 "at": datetime.now(timezone.utc).isoformat()})
        plan = await run_plan(api_client, store, batch, ladder=ladder, progress=progress)
    except Exception as exc:  # noqa: BLE001 — the reason goes on the row for the page to show
        log.warning("import plan %s failed: %s", batch_id, exc)
        await store.finish_plan(batch_id, plan=None, status=PLAN_FAILED, error=f"{exc}"[:2000])
        raise
    await store.finish_plan(batch_id, plan=plan, status=PLAN_READY)
    return plan


def plan_payload(batch: dict[str, Any]) -> dict[str, Any]:
    """What ``GET /prospects/import/{batchId}`` answers while the worker plans
    and once it has: the status, the progress, then the plan."""
    status = batch.get("plan_status") or (PLAN_READY if batch.get("plan", {}).get("rows") is not None else PLAN_QUEUED)
    out: dict[str, Any] = {
        "batchId": batch["id"], "name": batch["name"], "source": batch["door"], "planStatus": status,
        "planError": batch.get("plan_error"), "progress": batch.get("plan_progress"),
        "status": batch.get("status"), "fileName": batch.get("file_name"),
        "templateId": batch.get("template_id"), "templateName": (batch.get("plan") or {}).get("templateName") or "",
        "createdAt": batch["created_at"].isoformat() if batch.get("created_at") else None,
        "appliedAt": batch["applied_at"].isoformat() if batch.get("applied_at") else None,
        "undoneAt": batch["undone_at"].isoformat() if batch.get("undone_at") else None,
        "undoneBy": batch.get("undone_by"), "undoResult": batch.get("undo_result"),
        "result": batch.get("result"),
    }
    if status == PLAN_READY:
        out.update(_plan_answer(batch, batch.get("plan") or {}))
    return out


async def run_plan_cycle(settings: Settings, *, store=None, client: Optional[EspoApi] = None) -> int:
    """The worker's pass (ruling 123): claim one queued plan and run it. Returns
    how many were taken (0 or 1). Best-effort: never raises."""
    if not (settings.campaigns_enabled and settings.import_templates and settings.store_enabled):
        return 0
    from .ladder import _client as ladder_client, _store as ladder_store, available as ladder_available
    try:
        store = store if store is not None else ladder_store(settings)
        client = client or ladder_client(settings)
        if store is None or client is None:
            return 0
        batch = await store.claim_plan()
        if not batch:
            return 0
        ladder = bool(settings.lead_ladder and await ladder_available(client))
        try:
            plan = await plan_batch_now(client, store, batch["id"], ladder=ladder)
            log.info("import plan %s (%s): %s", batch["id"], batch.get("name"), plan.get("counts"))
        except Exception:  # noqa: BLE001 — recorded on the row by plan_batch_now
            pass
        return 1
    except Exception as exc:  # noqa: BLE001 — a cycle failure is a log line, never a crash
        log.warning("import plan cycle failed: %s", exc)
        return 0


# --- apply ---------------------------------------------------------------------------------

def entry_status(source: str, row: dict[str, Any]) -> str:
    """Ruling 53: a state filing with a website or an email enters Qualified on
    the outside filter's verdict; every other row enters Unqualified."""
    if source == cfg.SOURCE_FILING and (row.get("website") or row.get("email")):
        return cfg.LEAD_QUALIFIED
    return cfg.LEAD_UNQUALIFIED


def entry_status_for(settings: dict[str, Any], row: dict[str, Any]) -> str:
    """The status a row enters at. A template's *Enters at* (ruling 111)
    replaces ruling 53's fixed rule — with ruling 42's floor kept: a row with
    neither a website nor an email can never be Qualified, whatever the
    template says. Without a template, ruling 53 as before."""
    enters_at = settings.get("entersAt")
    if not enters_at:
        return entry_status(settings["door"], row)
    if enters_at == cfg.LEAD_QUALIFIED and (row.get("website") or row.get("email")):
        return cfg.LEAD_QUALIFIED
    return cfg.LEAD_UNQUALIFIED


def _lead_payload(row: dict[str, Any], plan: dict[str, Any], batch_name: str,
                  entry: dict[str, Any], *, ladder_fields: bool = False) -> dict[str, Any]:
    """The lead a row becomes. With ``ladder_fields`` (the CRM has the
    redesign's fields and the switch is on) a Qualified row carries the level
    its facts earn (build plan § 7e.3) — computed here, so the import at volume
    makes no extra CRM call per row — and the template's referrer partner."""
    source = plan["door"]
    status = entry_status_for(plan, row)
    now = _now_crm()
    payload: dict[str, Any] = {
        "name": (row.get("business_name") or f"{row.get('contact_first_name','')} {row.get('contact_last_name','')}".strip())[:255],
        "businessName": (row.get("business_name") or "")[:255],
        "firstName": (row.get("contact_first_name") or "")[:100] or None,
        "lastName": (row.get("contact_last_name") or "")[:100] or None,
        "mailingAddressStreet": row.get("mailing_street") or None,
        "mailingAddressCity": row.get("mailing_city") or None,
        "mailingAddressState": row.get("mailing_state") or None,
        "mailingAddressPostalCode": row.get("mailing_zip") or None,
        "mailingAddressCountry": "US",
        "addressStatus": cfg.ADDRESS_USABLE,
        "website": row.get("website") or None,
        "email": row.get("email") or None,
        "phone": row.get("phone") or None,
        "leadStatus": status,
        "leadStatusChangedAt": now,
        "leadSource": source,
        "externalReference": (row.get("external_reference") or "")[:100] or None,
        "importBatch": (batch_name if source == cfg.SOURCE_FILING
                        else f"{plan.get('listSource') or 'list'} — {batch_name}")[:255],
        "importRow": json.dumps(row, ensure_ascii=False)[:60000],
    }
    if ladder_fields and plan.get("referrerPartnerId"):
        payload["referredByPartnerId"] = plan["referrerPartnerId"]
    if status == cfg.LEAD_QUALIFIED:
        payload["qualifiedBy"] = (row.get("qualified_by") or plan.get("qualifiedBy") or "")[:255] or None
        payload["qualifiedAt"] = now
        if ladder_fields:
            from .ladder import facts, placed
            payload["qualificationLevel"] = placed(facts(payload))
    if entry.get("outcome") == OUTCOME_ALREADY_KNOWN:
        payload["sideExit"] = cfg.EXIT_DUPLICATE
        payload["sideExitAt"] = now
        match = entry.get("match") or {}
        what = "the contact" if match.get("kind") == "contact" else "the company"
        payload["sideExitNote"] = (f"Duplicate of {what} {match.get('name') or match.get('id')} "
                                   f"(matched by {match.get('by')} at import).")
    return {k: v for k, v in payload.items() if v is not None}


async def _second_partner(api_client: EspoApi, lead_id: str, partner_id: str) -> dict[str, Any]:
    """The template's partner on a known lead, under the API key: it becomes
    the lead's referrer when it has none, else a second association. Returns
    nothing for the user's fill. Never raises."""
    try:
        cur = await api_client.get(cfg.PROSPECT, lead_id, select="id,referredByPartnerId")
    except EspoError:
        return {}
    have = (cur or {}).get("referredByPartnerId")
    if not have:
        # Under the API key, like the create: the referrer is the import's to set.
        try:
            await api_client.update(cfg.PROSPECT, lead_id, {"referredByPartnerId": partner_id})
        except EspoError:
            pass
        return {}
    if have != partner_id:
        from core.partner_links import REASON_REFERRED_TO, associate
        await associate(api_client, reason=REASON_REFERRED_TO, partner_id=partner_id, lead_id=lead_id)
    return {}


def _match_links(entry: dict[str, Any]) -> dict[str, str]:
    """A Duplicate's link to the Company it matched. A matched Contact is named
    in the reason note only — never on ``promotedContact`` (*Converted to
    contact*), so a duplicate never reads as converted."""
    match = entry.get("match") or {}
    if entry.get("outcome") != OUTCOME_ALREADY_KNOWN or not match.get("id"):
        return {}
    if match.get("kind") == "company":
        return {"alreadyKnownCompanyId": match["id"]}
    return {}


async def apply_import(api_client: EspoApi, user_client: EspoApi, store, batch: dict[str, Any], *,
                       user_name: str, max_rows: Optional[int] = None) -> dict[str, Any]:
    """Write exactly the planned batch. Resumable: a row the ledger already
    holds is skipped, so a failed apply is re-run, never doubled.

    ``max_rows`` writes at most that many rows and returns with ``remaining``
    (the browser calls again until it is 0): 800 creates in one request
    outlast every timeout between the browser and the CRM (10-20-26)."""
    if batch.get("status") == "applied":
        raise CampaignError("This batch has already been applied.")
    if batch.get("plan_status") not in (None, PLAN_READY):
        raise CampaignError("This import has not finished its preview yet; wait for it, then apply.")
    plan = batch.get("plan") or {}
    rows = plan.get("rows") or []
    excluded = set(batch.get("excluded_rows") or [])
    # A row the preview found pointing at a lead the CRM no longer has: its
    # stale ledger rows are dropped first, or the across-batch check below
    # would skip the create (round two of the M4 pass, 10-25-26, v0.286.3).
    stale = [e["hash"] for e in rows if e.get("stale") and e["hash"] not in excluded]
    if stale:
        await store.forget_hashes(stale)
    done = await store.known_row_hashes([e["hash"] for e in rows])
    # A fill's hash is known across batches by definition: for fills "done"
    # means recorded by THIS batch.
    done_here = await store.batch_row_hashes(batch["id"])
    result = {"created": 0, "duplicates": 0, "filled": 0, "qualified": 0, "unqualified": 0,
              "excluded": 0, "skipped": 0, "errors": 0}
    ledger: list[dict[str, Any]] = []
    writes = 0
    remaining = 0
    # The preview showed the levels, so the apply writes exactly what it showed
    # (ruling 83): a templated plan carries whether the ladder was on; the CSV
    # path asks the switch and the CRM as before.
    ladder_fields = False
    if plan.get("ladder") is not None:
        ladder_fields = bool(plan["ladder"])
    elif get_settings().lead_ladder:
        from .ladder import available as ladder_available
        ladder_fields = await ladder_available(api_client)
    for entry in rows:
        outcome = entry.get("outcome")
        if entry["hash"] in excluded and outcome in (OUTCOME_CREATE, OUTCOME_ALREADY_KNOWN, OUTCOME_FILLS):
            result["excluded"] += 1
            continue
        if outcome == OUTCOME_FILLS:
            if entry["hash"] in done_here or not entry.get("fill") or not entry.get("leadId"):
                result["skipped"] += 1
                continue
            if max_rows is not None and writes >= max_rows:
                remaining += 1
                continue
            writes += 1
            try:
                # Ruling 84: only the empty fields, re-checked against the lead
                # as it is now, written as the user, then placement.
                lead = await api_client.get(cfg.PROSPECT, entry["leadId"], select=FILL_SELECT)
                if not lead or not lead.get("id"):
                    result["skipped"] += 1
                    continue
                fill = {k: v for k, v in entry["fill"].items() if not str(lead.get(k) or "").strip()}
                # The template's partner on a lead already here: it becomes the
                # lead's referrer when it has none, else a second association
                # (prds/partner-association-plan.md, decision 2).
                if ladder_fields and plan.get("referrerPartnerId"):
                    fill.update(await _second_partner(api_client, entry["leadId"], plan["referrerPartnerId"]))
                if fill:
                    await user_client.update(cfg.PROSPECT, entry["leadId"], fill)
                    from .ladder import after_write
                    await after_write(entry["leadId"], client=api_client)
                result["filled"] += 1
                ledger.append({"row_hash": entry["hash"], "external_reference": entry["row"].get("external_reference") or None,
                               "business_name": entry.get("name"), "postal_code": entry.get("postalCode"),
                               "outcome": OUTCOME_FILLS, "prospect_id": entry["leadId"]})
                done_here.add(entry["hash"])
            except EspoError as exc:
                result["errors"] += 1
                log.warning("import %s fill %r failed: %s", batch["id"], entry.get("name"), exc)
            if len(ledger) >= 50:
                await store.record_rows(batch["id"], ledger)
                ledger = []
            continue
        if outcome not in (OUTCOME_CREATE, OUTCOME_ALREADY_KNOWN) or entry["hash"] in done:
            result["skipped"] += 1
            continue
        if max_rows is not None and writes >= max_rows:
            remaining += 1
            continue
        writes += 1
        row = entry["row"]
        try:
            payload = _lead_payload(row, plan, batch["name"], entry, ladder_fields=ladder_fields)
            # The lead is the signed-in user's (ruling 33). Its link to a
            # matched Company is the API client's second write: EspoCRM checks
            # read access on the foreign record when a link rides a create, and
            # the Marketing Admin Role reads no Company (found live 10-15-26).
            lead = await user_client.create(cfg.PROSPECT, payload)
            links = _match_links(entry)
            if links:
                await api_client.update(cfg.PROSPECT, lead["id"], links)
            if outcome == OUTCOME_CREATE:
                result["created"] += 1
                result["qualified" if payload["leadStatus"] == cfg.LEAD_QUALIFIED else "unqualified"] += 1
            else:
                result["duplicates"] += 1
            result.setdefault("firstProspectId", lead.get("id"))
            ledger.append({"row_hash": entry["hash"], "external_reference": row.get("external_reference") or None,
                           "business_name": payload.get("businessName"), "postal_code": payload.get("mailingAddressPostalCode"),
                           "outcome": outcome, "prospect_id": lead.get("id")})
            done.add(entry["hash"])
        except EspoError as exc:
            result["errors"] += 1
            log.warning("import %s row %r failed: %s", batch["id"], entry.get("name"), exc)
        if len(ledger) >= 50:
            await store.record_rows(batch["id"], ledger)
            ledger = []
    if ledger:
        await store.record_rows(batch["id"], ledger)
    result["appliedBy"] = user_name
    result["remaining"] = remaining
    if result["errors"] == 0 and remaining == 0:
        await store.mark_applied(batch["id"], result)
    return result


# --- Undo (rulings 87, 127; M4) -----------------------------------------------------------------

def _parse_crm_time(value: Any) -> Optional[datetime]:
    s = str(value or "").strip()
    if not s:
        return None
    try:
        return datetime.strptime(s[:19], "%Y-%m-%d %H:%M:%S").replace(tzinfo=timezone.utc)
    except ValueError:
        try:
            dt = datetime.fromisoformat(s)
            return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)
        except ValueError:
            return None


def touched_reasons(lead: dict[str, Any], written_at: Optional[datetime], *, has_touch: bool,
                    in_list: bool, has_conversation: bool, filled_by: str = "") -> list[str]:
    """Why a lead the batch created is *touched* (ruling 87) — every reason, in
    plain words; an empty list means Undo may delete it. A change inside
    :data:`UNDO_WINDOW` of the write is the import's own."""
    reasons: list[str] = []
    cutoff = (written_at + UNDO_WINDOW) if written_at else None

    def later(value: Any) -> bool:
        dt = _parse_crm_time(value)
        return bool(dt and (cutoff is None or dt > cutoff))

    if has_touch:
        reasons.append("has a touch")
    if in_list:
        reasons.append("is on a campaign list")
    if has_conversation:
        reasons.append("has a conversation")
    if filled_by:
        reasons.append(f"filled by a later import ({filled_by})")
    if later(lead.get("leadStatusChangedAt")):
        reasons.append(f"status changed since ({lead.get('leadStatus') or 'unknown'})")
    if str(lead.get("levelSetAt") or "").strip():
        reasons.append("level set by hand")
    exit_ = lead.get("sideExit")
    if exit_ and not (exit_ == cfg.EXIT_DUPLICATE and not later(lead.get("sideExitAt"))):
        reasons.append(f"stopped ({exit_})")
    if later(lead.get("modifiedAt")) and not reasons:
        reasons.append("edited since the import")
    return reasons


async def _conversations_available(api_client: EspoApi) -> bool:
    reader = getattr(api_client, "metadata", None)
    if reader is None:
        return False
    try:
        links = await reader(f"entityDefs.{cfg.PROSPECT}.links")
    except Exception:  # noqa: BLE001 — unreadable is "no such link"
        return False
    return isinstance(links, dict) and CONVERSATIONS_LINK in links


async def _conversed(api_client: EspoApi, ids: list[str], available: bool) -> set[str]:
    """The leads among ``ids`` with a conversation, feature-detected. One read
    asks whether the chunk has any at all; only then one read per lead."""
    if not available or not ids:
        return set()
    try:
        page = await api_client.list("CConversation", select="id", max_size=1,
                                     where=[{"type": "linkedWith", "attribute": "prospects", "value": ids}])
        if not page.get("total") and not page.get("list"):
            return set()
    except EspoError as exc:
        log.warning("undo: conversation check for %s leads failed: %s", len(ids), exc)
        return set()
    out: set[str] = set()
    for lead_id in ids:
        try:
            page = await api_client.list_related(cfg.PROSPECT, lead_id, CONVERSATIONS_LINK, select="id", max_size=1)
            if page.get("list"):
                out.add(lead_id)
        except EspoError as exc:
            log.warning("undo: conversations of %s unreadable (%s) — counted as touched", lead_id, exc)
            out.add(lead_id)
    return out


async def undo_import(api_client: EspoApi, store, batch: dict[str, Any], *, actor: str,
                      max_leads: Optional[int] = None) -> dict[str, Any]:
    """*Undo this import* (ruling 87): delete every lead the batch created
    that nobody has touched since, under the org-wide key (ruling 127); list
    the touched ones and leave them; fills stay. In parts of ``max_leads``,
    like the apply — the page calls again until ``remaining`` is 0, and the
    batch is stamped undone on the last part."""
    if batch.get("status") != "applied":
        raise CampaignError("Only an applied import can be undone.")
    if batch.get("undone_at"):
        raise CampaignError("This import has already been undone.")
    created = await store.batch_created_leads(batch["id"])
    # A touched lead stays in the ledger, so an earlier part's touched list is
    # what lets the next part move on past it.
    already = {t.get("id") for t in (batch.get("undo_result") or {}).get("touched") or []}
    pending = [r for r in created if r["prospect_id"] not in already]
    todo = pending if max_leads is None else pending[:max_leads]
    remaining = len(pending) - len(todo)
    deleted: list[str] = []
    touched: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    conversations = await _conversations_available(api_client)
    for i in range(0, len(todo), _ID_CHUNK):
        chunk = todo[i:i + _ID_CHUNK]
        ids = [r["prospect_id"] for r in chunk]
        written = {r["prospect_id"]: r.get("created_at") for r in chunk}
        names = {r["prospect_id"]: r.get("business_name") or r["prospect_id"] for r in chunk}
        leads: dict[str, dict[str, Any]] = {}
        try:
            page = await api_client.list(cfg.PROSPECT, select=UNDO_SELECT, max_size=len(ids),
                                         where=[{"type": "in", "attribute": "id", "value": ids}])
            leads = {l["id"]: l for l in page.get("list", [])}
        except EspoError as exc:
            errors.append({"ids": ids, "error": f"could not read the leads: {exc}"[:500]})
            continue
        with_touch: set[str] = set()
        try:
            page = await api_client.list(cfg.TOUCH, select="id,prospectId", max_size=200,
                                         where=[{"type": "in", "attribute": "prospectId", "value": ids}])
            with_touch = {t.get("prospectId") for t in page.get("list", [])}
        except EspoError as exc:
            log.warning("undo: touch check for %s leads failed (%s) — counted as touched", len(ids), exc)
            with_touch = set(ids)
        try:
            listed = {k.split(":", 1)[1] for k in await store.list_memberships([f"lead:{i}" for i in ids])}
        except Exception as exc:  # noqa: BLE001 — an unreadable ledger counts as touched
            log.warning("undo: list membership check failed (%s) — counted as touched", exc)
            listed = set(ids)
        conversed = await _conversed(api_client, ids, conversations)
        try:
            filled = await store.filled_by_later_batch(batch["id"], ids)
        except Exception as exc:  # noqa: BLE001 — an unreadable ledger counts as touched
            log.warning("undo: later-fill check failed (%s) — counted as touched", exc)
            filled = {i: "an import the ledger could not read" for i in ids}
        gone: list[str] = []
        chunk_deleted: list[str] = []
        for lead_id in ids:
            lead = leads.get(lead_id)
            if not lead:
                gone.append(lead_id)           # already deleted: nothing to undo, nothing to keep
                continue
            reasons = touched_reasons(lead, written.get(lead_id), has_touch=lead_id in with_touch,
                                      in_list=lead_id in listed, has_conversation=lead_id in conversed,
                                      filled_by=filled.get(lead_id, ""))
            if reasons:
                touched.append({"id": lead_id, "name": lead.get("businessName") or lead.get("name") or names[lead_id],
                                "why": "; ".join(reasons)})
                continue
            try:
                await api_client.delete(cfg.PROSPECT, lead_id)
                deleted.append(lead_id)
                chunk_deleted.append(lead_id)
            except EspoError as exc:
                if is_not_found(exc):
                    gone.append(lead_id)
                else:
                    errors.append({"id": lead_id, "name": names[lead_id], "error": f"{exc}"[:500]})
        await store.forget_rows(batch["id"], chunk_deleted + gone)
    done = remaining == 0 and not errors
    result = await store.record_undo(batch["id"], deleted=deleted, touched=touched, errors=errors,
                                     done=done, actor=actor)
    return {"deleted": deleted, "touched": touched, "errors": errors, "remaining": remaining,
            "done": done, "total": result}
